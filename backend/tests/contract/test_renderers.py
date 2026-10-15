"""Four migrated builder façades still generate valid scoped documents."""
import io
import uuid
import zipfile

import pytest

pytestmark = [pytest.mark.integration, pytest.mark.asyncio(loop_scope="session")]


async def test_pegawai_xlsx_retains_columns_nip_text_and_satker_scope(contract):
    await contract.db.pegawai.insert_many([
        {"id": "renderer-employee-a", "nama": "Pegawai render A", "nip": "001234567890123456", "kode_satker": contract.users["a"]["kode_satker"]},
        {"id": "renderer-employee-b", "nama": "Pegawai rahasia B", "nip": "009999999999999999", "kode_satker": contract.users["b"]["kode_satker"]},
    ])
    response = await contract.http.get("/api/pegawai/export-xlsx", headers=contract.headers["a"])
    assert response.status_code == 200, response.text[:150]
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Data Pegawai", "Referensi"]
    rows = list(workbook["Data Pegawai"].values)
    assert any("001234567890123456" in row and "Pegawai render A" in row for row in rows)
    assert all("Pegawai rahasia B" not in row for row in rows)
    assert workbook["Data Pegawai"].freeze_panes == "C2"


async def test_kodefikasi_snapshot_both_formats_preserve_hierarchy(contract):
    await contract.db.kodefikasi.update_one({"kode": "3999901001"}, {"$set": {
        "kode": "3999901001", "level": 5, "parent_kode": "3999901", "uraian": "Kode snapshot sintetis", "satuan": "Unit"}}, upsert=True)
    for shape in ("datar", "hierarki"):
        response = await contract.http.get("/api/kodefikasi/export", params={"bentuk": shape}, headers=contract.headers["a"])
        assert response.status_code == 200, response.text[:150]
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        rows = list(workbook["Kodefikasi"].values)
        assert any("3999901001" in row and "Kode snapshot sintetis" in row for row in rows)
        if shape == "hierarki":
            row = next(row for row in rows if "3999901001" in row)
            assert row[0] == "3" and row[8] == "3999901001" and row[10] == "Unit"


async def test_stiker_pdf_remains_valid_after_thread_extraction(contract):
    aid = "renderer-stiker-" + uuid.uuid4().hex
    await contract.db.assets.insert_one({"id": aid, "activity_id": "contract-activity-a",
        "asset_name": "Stiker sintetis", "asset_code": "3050104001", "NUP": "1", "purchase_price": "1000"})
    response = await contract.http.get("/api/stiker/label", params={"asset_ids": aid, "sampel_ukuran": "false"}, headers=contract.headers["a"])
    assert response.status_code == 200, response.text[:150]
    assert response.content.startswith(b"%PDF-") and response.headers["X-Total-Stiker"] == "1"
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(response.content))
    assert len(reader.pages) >= 1 and "Stiker sintetis" in "".join(page.extract_text() for page in reader.pages)


async def test_lbp_docx_async_data_and_thread_document_builder(contract):
    await contract.db.assets.insert_one({"id": "renderer-lbp-foreign", "activity_id": "contract-activity-b",
        "asset_name": "Aset rahasia B", "asset_code": "3050104001", "NUP": "9", "purchase_price": "1234567891", "purchase_date": "2026-01-01"})
    response = await contract.http.get("/api/pelaporan/lbp-docx", params={"tahun": 2026, "semester": 2}, headers=contract.headers["a"])
    assert response.status_code == 200, response.text[:150]
    with zipfile.ZipFile(io.BytesIO(response.content)) as document:
        content = document.read("word/document.xml").decode()
        assert "LAPORAN BARANG PENGGUNA" in content and "KATA PENGANTAR" in content
        assert "Aset rahasia B" not in content and "1.234.567.891" not in content


@pytest.mark.parametrize("count", [1, 2, 3])
async def test_pegawai_limit_is_explicit_instead_of_truncating(contract, monkeypatch, count):
    import routes.pegawai as pegawai
    monkeypatch.setattr(pegawai, "MAX_PEGAWAI_EXPORT", 2)
    code = contract.users["b"]["kode_satker"]
    await contract.db.pegawai.delete_many({"kode_satker": code})
    await contract.db.pegawai.insert_many([{"id": f"cap-{n}", "nama": f"Cap {n}", "kode_satker": code} for n in range(count)])
    response = await contract.http.get("/api/pegawai/export-xlsx", headers=contract.headers["b"])
    assert response.status_code == (413 if count > 2 else 200)


async def test_kodefikasi_snapshot_byte_and_row_limits_fail_explicitly(contract, monkeypatch):
    import routes.kodefikasi as kodefikasi
    monkeypatch.setattr(kodefikasi, "MAX_KODE_EXPORT_ROWS", 0)
    response = await contract.http.get("/api/kodefikasi/export", headers=contract.headers["a"])
    assert response.status_code == 413
    monkeypatch.setattr(kodefikasi, "MAX_KODE_EXPORT_ROWS", 250000)
    monkeypatch.setattr(kodefikasi, "MAX_KODE_SNAPSHOT_BYTES", 1)
    response = await contract.http.get("/api/kodefikasi/export", headers=contract.headers["a"])
    assert response.status_code == 413
