"""Label operasional konsisten pada keluaran; data historis tidak ditulis ulang."""
import asyncio
import csv
import io
import os
from pathlib import Path

import openpyxl
import pytest
from mongomock_motor import AsyncMongoMockClient
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from starlette.requests import Request

from operasional_utils import JENIS_OPERASIONAL, nama_jenis_operasional


@pytest.mark.parametrize("nilai,harapan", [
    ("Kegiatan/Acara/Kebutuhan", "Unit/Tempat/Tugas"),
    (" kegiatan / acara / kebutuhan ", "Unit/Tempat/Tugas"),
    ("Unit/Tempat/Tugas", "Unit/Tempat/Tugas"),
    (" unit / tempat / tugas ", "Unit/Tempat/Tugas"),
    ("Ruangan", "Ruangan"), (" ruangan ", "Ruangan"),
    ("Khusus", "Khusus"), ("", ""), (None, ""),
])
def test_nama_lama_tetap_dikenali(nilai, harapan):
    assert nama_jenis_operasional(nilai) == harapan


def test_template_impor_memakai_referensi_baru():
    from routes.templates import ASSET_TEMPLATE_SCHEMA
    field = next(f for f in ASSET_TEMPLATE_SCHEMA if f["field"] == "operasional_jenis")
    assert field["dropdown"] == list(JENIS_OPERASIONAL)
    assert "Unit/Tempat/Tugas" in field["rule"]
    assert "Kegiatan/Acara/Kebutuhan" not in field["rule"]


def test_csv_dan_xlsx_menampilkan_alias_tanpa_mengubah_database(monkeypatch):
    import routes.exports as ex
    import shared_utils as su
    database = AsyncMongoMockClient()["uji_operasional"]
    monkeypatch.setattr(ex, "db", database)
    monkeypatch.setattr(su, "db", database)

    async def jalankan():
        await database.assets.insert_many([
            {"id": "a1", "asset_code": "3050105007", "asset_name": "Kursi", "NUP": "1",
             "pengguna_melekat_ke": "Operasional", "operasional_jenis": "Kegiatan/Acara/Kebutuhan"},
            {"id": "a2", "asset_name": "Meja", "operasional_jenis": "Ruangan"},
        ])
        user = {"id": "admin-uji", "role": "admin", "kode_satker": ""}
        request = Request({"type": "http", "method": "GET", "path": "/api/export/csv", "headers": []})
        response = await ex.export_csv.__wrapped__(request, activity_id=None, filter_aset={}, _user=user)
        bagian = [p.decode() if isinstance(p, bytes) else p async for p in response.body_iterator]
        rows = list(csv.DictReader(io.StringIO("".join(bagian))))
        assert [r["operasional_jenis"] for r in rows] == ["Unit/Tempat/Tugas", "Ruangan"]
        data = await ex.bangun_xlsx_bytes({})
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        try:
            sheet = workbook["Data Aset"]
            # Kolom W = jenis operasional, sesuai registry/header ekspor.
            assert sheet.cell(2, 23).value == "Unit/Tempat/Tugas"
            assert sheet.cell(3, 23).value == "Ruangan"
        finally:
            workbook.close()
        assert (await database.assets.find_one({"id": "a1"}))["operasional_jenis"] == "Kegiatan/Acara/Kebutuhan"
    asyncio.run(jalankan())


def test_kartu_pdf_memakai_nama_baru_tanpa_merusak_aset():
    from routes.cards import _gambar_kartu_ke_kanvas
    aset = {"id": "a1", "asset_code": "3050105007", "asset_name": "Kursi Uji", "NUP": "1",
            "pengguna_melekat_ke": "Operasional", "operasional_jenis": "Kegiatan/Acara/Kebutuhan",
            "user": "Penanggung Jawab Uji"}
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=landscape(A4))
    _gambar_kartu_ke_kanvas(pdf, [(aset, [])])
    pdf.save()
    data = buffer.getvalue()
    teks = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)
    assert "Unit/Tempat/Tugas" in teks
    assert "Kegiatan/Acara/Kebutuhan" not in teks
    assert aset["operasional_jenis"] == "Kegiatan/Acara/Kebutuhan"
    if os.environ.get("AMAN_QA_PDF_DIR"):
        folder = Path(os.environ["AMAN_QA_PDF_DIR"])
        folder.mkdir(parents=True, exist_ok=True)
        (folder / "kartu-jenis-operasional.pdf").write_bytes(data)
