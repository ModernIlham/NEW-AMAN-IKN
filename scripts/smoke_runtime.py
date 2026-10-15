#!/usr/bin/env python3
"""Prove a runtime-only install can authenticate and render actual API outputs.

Always creates/drops its own aman_ci_runtime_* database on explicit loopback.
No pytest/dev dependency, preexisting account, application .env or external API.
"""
import argparse
import asyncio
import hashlib
import importlib.metadata
import io
import json
import os
from pathlib import Path
import secrets
import sys
import tempfile
import uuid


async def smoke(args):
    root = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(root / "backend"))
    from fixture_target import mongo_loopback
    args.mongo_url = mongo_loopback(args.mongo_url)
    database = "aman_ci_runtime_" + uuid.uuid4().hex
    os.environ.update(MONGO_URL=args.mongo_url, DB_NAME=database,
                      JWT_SECRET=secrets.token_urlsafe(48),
                      ADMIN_BOOTSTRAP_TOKEN=secrets.token_urlsafe(48),
                      REDIS_URL="", MEILI_URL="", MEILI_MASTER_KEY="",
                      TINIFY_API_KEY="", RESEND_API_KEY="")
    root = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(root / "backend"))
    # Imports deliberately occur after environment isolation and within the
    # running loop; Motor/GridFS must never bind to a different event loop.
    import httpx
    import slowapi
    original_limiter = slowapi.Limiter
    def isolated_limiter(*positional, **options):
        # The application normally uses a dedicated aman_ratelimit DB. This
        # probe must write ONLY its generated DB, so retain real rate limiting
        # with local MemoryStorage. Distributed limiter tests live in contract.
        options["storage_uri"] = "memory://"
        options.pop("storage_options", None)
        return original_limiter(*positional, **options)
    slowapi.Limiter = isolated_limiter
    from db import client, db
    from server import app
    slowapi.Limiter = original_limiter
    import routes.backup as backup
    sandbox = tempfile.TemporaryDirectory(prefix="aman-runtime-private-")
    previous_paths = backup.BACKUP_TEMP_DIR, backup.BACKUP_ARSIP_DIR, backup.UPLOADS_DIR
    for name, directory in (("BACKUP_TEMP_DIR", "backup-temp"), ("BACKUP_ARSIP_DIR", "backup-arsip"), ("UPLOADS_DIR", "uploads")):
        path = Path(sandbox.name) / directory
        path.mkdir(mode=0o700)
        setattr(backup, name, path)
    from indexes import create_indexes, kegagalan_indeks
    from docx import Document
    from pptx import Presentation
    from pypdf import PdfReader
    from openpyxl import load_workbook
    from weasyprint import HTML
    from scripts.synthdata import generate_assets
    from backup_archive import validate_archive
    import zipfile

    assert db.name == database and database.startswith("aman_ci_runtime_")
    result = {"python": sys.version.split()[0], "runtime_only": True,
              "rate_limit_storage": "isolated-memory", "checks": {}}
    lifespan_task = None
    lifecycle_receive, lifecycle_send = asyncio.Queue(), asyncio.Queue()
    try:
        if args.lifespan:
            lifespan_task = asyncio.create_task(app(
                {"type": "lifespan", "asgi": {"version": "3.0", "spec_version": "2.0"}},
                lifecycle_receive.get, lifecycle_send.put))
            await lifecycle_receive.put({"type": "lifespan.startup"})
            message = await asyncio.wait_for(lifecycle_send.get(), 60)
            assert message["type"] == "lifespan.startup.complete", "ASGI startup failed"
            result["checks"]["asgi_startup"] = True
        await db.command("ping")
        await create_indexes()
        assert not kegagalan_indeks(), "Runtime index creation failed"
        password = "Runtime9-" + secrets.token_urlsafe(20)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),
                                     base_url="http://runtime-fixture", timeout=60) as http:
            bootstrap = await http.post("/api/auth/bootstrap", json={
                "username": "runtime-fixture", "password": password, "name": "Runtime fixture"},
                headers={"X-Admin-Bootstrap-Token": os.environ["ADMIN_BOOTSTRAP_TOKEN"]})
            assert bootstrap.status_code == 200, "Bootstrap runtime failed"
            login = await http.post("/api/auth/login", json={"username": "runtime-fixture", "password": password})
            assert login.status_code == 200, "Login runtime failed"
            headers = {"Authorization": "Bearer " + login.json()["access_token"]}
            assert (await http.get("/api/auth/me", headers=headers)).status_code == 200
            result["checks"]["auth_bootstrap_login_me"] = True
            activity = "runtime-fixture-activity"
            await db.inventory_activities.insert_one({"id": activity, "nama_kegiatan": "Runtime fixture", "kode_satker": ""})
            rows = generate_assets(3)
            for index, row in enumerate(rows):
                row.update(id=f"runtime-asset-{index}", activity_id=activity, kode_satker="")
            await db.assets.insert_many(rows)
            await db.categories.insert_one({"id": "runtime-category", "label": "Runtime fixture"})
            for kind, path in (("pdf", "/export/pdf"), ("xlsx", "/export/xlsx"),
                               ("pptx", "/documents/ppt"), ("docx", "/documents/proposal")):
                response = await http.get("/api" + path, headers=headers,
                                          params={"activity_id": activity})
                assert response.status_code == 200, f"{kind} API status {response.status_code}"
                data = response.content
                if kind == "pdf":
                    document = PdfReader(io.BytesIO(data))
                    assert len(document.pages) > 0
                    geometry = [[float(page.mediabox.width), float(page.mediabox.height)] for page in document.pages]
                elif kind == "xlsx":
                    workbook = load_workbook(io.BytesIO(data))
                    assert workbook.active.max_row > 3
                    geometry = {sheet.title: {"rows": sheet.max_row, "columns": sheet.max_column}
                                for sheet in workbook}
                elif kind == "pptx":
                    presentation = Presentation(io.BytesIO(data))
                    assert len(presentation.slides) > 0
                    geometry = {"slides": len(presentation.slides), "width_emu": presentation.slide_width,
                                "height_emu": presentation.slide_height}
                else:
                    document = Document(io.BytesIO(data))
                    assert document.paragraphs or document.tables
                    geometry = [{"width_emu": section.page_width, "height_emu": section.page_height}
                                for section in document.sections]
                result["checks"][kind] = {"bytes": len(data), "geometry": geometry}
            native_pdf = HTML(string="<style>@page{size:A4;margin:20mm}</style><p>AMAN runtime Pango Inter</p>").write_pdf()
            assert len(PdfReader(io.BytesIO(native_pdf)).pages) == 1
            result["checks"]["weasyprint_native_pango"] = True
            # Exercise the actual archive builder and manifest validator too.
            import routes.backup as backup
            from tempfile import TemporaryDirectory
            with TemporaryDirectory(prefix="aman-runtime-backup-") as temp:
                previous = (backup.BACKUP_TEMP_DIR, backup.BACKUP_ARSIP_DIR, backup.UPLOADS_DIR)
                try:
                    backup.BACKUP_TEMP_DIR = Path(temp) / "temp"
                    backup.BACKUP_ARSIP_DIR = Path(temp) / "archive"
                    backup.UPLOADS_DIR = Path(temp) / "uploads"
                    for path in (backup.BACKUP_TEMP_DIR, backup.BACKUP_ARSIP_DIR, backup.UPLOADS_DIR):
                        path.mkdir()
                    await db.backup_jobs.insert_one({"job_id": "runtime-backup", "type": "backup", "status": "queued"})
                    await backup.run_backup_task("runtime-backup", "runtime-fixture")
                    job = await db.backup_jobs.find_one({"job_id": "runtime-backup"})
                    assert job["status"] == "completed" and job["download_ready"] is True
                    with zipfile.ZipFile(backup.BACKUP_TEMP_DIR / "runtime-backup.zip") as archive:
                        checked = validate_archive(archive)
                    assert checked["collections"]["assets"] == 3
                    result["checks"]["backup_manifest"] = True
                finally:
                    backup.BACKUP_TEMP_DIR, backup.BACKUP_ARSIP_DIR, backup.UPLOADS_DIR = previous
        packages = sorted((dist.metadata["Name"].lower().replace("_", "-"), dist.version)
                          for dist in importlib.metadata.distributions()
                          if dist.metadata["Name"].lower() not in {"pip", "setuptools", "wheel"})
        result["distribution_count"] = len(packages)
        result["distribution_manifest_sha256"] = hashlib.sha256(json.dumps(packages).encode()).hexdigest()
        return result
    finally:
        try:
            if lifespan_task and not lifespan_task.done():
                await lifecycle_receive.put({"type": "lifespan.shutdown"})
                message = await asyncio.wait_for(lifecycle_send.get(), 30)
                assert message["type"] == "lifespan.shutdown.complete", "ASGI shutdown failed"
                await asyncio.wait_for(lifespan_task, 30)
                result["checks"]["asgi_shutdown"] = True
        finally:
            if lifespan_task and not lifespan_task.done():
                lifespan_task.cancel()
            assert db.name == database and database.startswith("aman_ci_runtime_")
            # Real application shutdown closes its client. A separate local
            # cleanup client owns only this generated database after teardown.
            from motor.motor_asyncio import AsyncIOMotorClient
            cleanup = AsyncIOMotorClient(args.mongo_url, serverSelectionTimeoutMS=5000)
            try:
                await cleanup.drop_database(database)
                assert database not in await cleanup.list_database_names()
            finally:
                cleanup.close()
                client.close()
                backup.BACKUP_TEMP_DIR, backup.BACKUP_ARSIP_DIR, backup.UPLOADS_DIR = previous_paths
                sandbox.cleanup()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mongo-url", required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--lifespan", action="store_true", help="Exercise actual ASGI startup/shutdown hooks")
    args = parser.parse_args()
    report = asyncio.run(smoke(args))
    rendered = json.dumps(report, indent=2, sort_keys=True) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered)
    else:
        print(rendered, end="")
