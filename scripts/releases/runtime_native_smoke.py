#!/usr/bin/env python3
"""DB-free native renderer preflight. All documents are synthetic in-memory."""
import argparse
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile

INTER_SHA256="29160a80ff49ddcab2c97711247e08b1fab27a484a329ce8b813d820dc559031"


def smoke(font_path):
    # Native font caches must not rely on a writable service HOME. No rendered
    # user documents are written; the temporary cache is private and bounded.
    with tempfile.TemporaryDirectory(prefix="aman-native-") as cache:
        previous=os.environ.get("XDG_CACHE_HOME")
        os.environ["XDG_CACHE_HOME"]=cache
        try:return _smoke(font_path)
        finally:
            if previous is None:os.environ.pop("XDG_CACHE_HOME",None)
            else:os.environ["XDG_CACHE_HOME"]=previous


def _smoke(font_path):
    font=Path(font_path)
    if font.is_symlink() or not font.is_file() or font.stat().st_size>2*1024*1024:
        raise ValueError("Pinned Inter font must be a bounded regular file")
    if hashlib.sha256(font.read_bytes()).hexdigest()!=INTER_SHA256:
        raise ValueError("Pinned Inter font checksum differs")
    from fontTools.ttLib import TTFont as SourceFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.pagesizes import A4
    from pypdf import PdfReader
    from weasyprint import HTML
    from weasyprint.text.fonts import FontConfiguration
    from openpyxl import Workbook,load_workbook
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    text="AMAN 0123456789"
    pdfmetrics.registerFont(TTFont("AMANDeployInter",str(font),validate=1))
    width=pdfmetrics.stringWidth(text,"AMANDeployInter",12)
    with SourceFont(font) as reference:
        family=reference["name"].getBestFamilyName()
        if "Inter" not in family:raise ValueError("Font family is not Inter")
        cmap=reference.getBestCmap()
        expected=sum(reference["hmtx"].metrics[cmap[ord(character)]][0] for character in text)/reference["head"].unitsPerEm*12
    if abs(width-expected)>.01:raise ValueError("Inter text geometry changed")
    report=BytesIO();canvas=Canvas(report,pagesize=A4)
    canvas.setFont("AMANDeployInter",12);canvas.drawString(40,800,text);canvas.save()
    if len(PdfReader(BytesIO(report.getvalue())).pages)!=1:raise ValueError("ReportLab PDF invalid")
    allowed_url=font.resolve().as_uri()
    try:
        from weasyprint.urls import URLFetcher,URLFetcherResponse,FatalURLFetchingError
        class PinnedFontFetcher(URLFetcher):
            def fetch(self,url,headers=None):
                if url!=allowed_url:raise FatalURLFetchingError("External resource forbidden in native smoke")
                return URLFetcherResponse(url,body=BytesIO(font.read_bytes()),headers={"Content-Type":"font/ttf"})
        fetcher=PinnedFontFetcher(allowed_protocols={"file"},allow_redirects=False,fail_on_errors=True)
    except ImportError:
        # Retained legacy cohorts use the earlier callable protocol.
        def fetcher(url,*args,**kwargs):
            if url!=allowed_url:raise ValueError("External resource forbidden in native smoke")
            return {"string":font.read_bytes(),"mime_type":"font/ttf"}
    html=('<style>@font-face{font-family:Inter;src:url("'+allowed_url+'")}'+
          '@page{size:A4;margin:20mm}body{font-family:Inter}</style><p>'+text+'</p>')
    pdf=HTML(string=html,url_fetcher=fetcher).write_pdf(font_config=FontConfiguration())
    reader=PdfReader(BytesIO(pdf))
    if len(reader.pages)!=1:raise ValueError("WeasyPrint page count changed")
    page=reader.pages[0]
    pdf_fonts=page["/Resources"].get("/Font",{})
    if not any("Inter" in str(value.get_object().get("/BaseFont","")) for value in pdf_fonts.values()):
        raise ValueError("Pinned Inter was not embedded; fallback font forbidden")
    geometry=[float(page.mediabox.width),float(page.mediabox.height)]
    if abs(geometry[0]-A4[0])>.1 or abs(geometry[1]-A4[1])>.1:raise ValueError("A4 PDF geometry changed")
    workbook=Workbook();workbook.active.append(["AMAN",123]);xlsx=BytesIO();workbook.save(xlsx)
    if load_workbook(BytesIO(xlsx.getvalue()),read_only=True).active["B1"].value!=123:raise ValueError("XLSX roundtrip failed")
    document=Document();document.add_paragraph("AMAN");docx=BytesIO();document.save(docx)
    if Document(BytesIO(docx.getvalue())).paragraphs[0].text!="AMAN":raise ValueError("DOCX roundtrip failed")
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(1)).text="AMAN"
    pptx=BytesIO();presentation.save(pptx)
    reopened=Presentation(BytesIO(pptx.getvalue()))
    if len(reopened.slides)!=1 or reopened.slides[0].shapes[0].text!="AMAN":raise ValueError("PPTX roundtrip failed")
    blobs={"reportlab_pdf":report.getvalue(),"weasy_pdf":pdf,"xlsx":xlsx.getvalue(),"docx":docx.getvalue(),"pptx":pptx.getvalue()}
    return {"ok":True,"network_resources":"denied_except_pinned_local_font","font_sha256":INTER_SHA256,
        "inter_width_points":round(width,6),"a4_points":geometry,"pdf_pages":1,"xlsx_rows":1,"docx_paragraphs":1,"pptx_slides":1,
        "artifacts":{name:{"bytes":len(blob),"sha256":hashlib.sha256(blob).hexdigest()} for name,blob in blobs.items()}}


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--font",required=True)
    args=parser.parse_args()
    try:print(json.dumps(smoke(args.font),sort_keys=True))
    except Exception:parser.exit(1,"Native renderer/font preflight failed; no application data was accessed.\n")
